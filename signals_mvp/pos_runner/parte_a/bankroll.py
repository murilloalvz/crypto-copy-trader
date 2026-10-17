"""Parte A do addendum de relatorio (2026-10-16, emendado no mesmo dia para a grade de 24 combinacoes): tabela
de trades + gestao de banca por combinacao (tf x saida x medida da queda x faixa de idade), loteada sobre todo um
periodo (discovery/holdout). A gestao de banca em si (`BankrollLedger`/`BankrollState`) vive em `..bankroll` (sem
dependencia de `Context`/`Store`) e e compartilhada com a Parte B ao vivo (`live.py`), que chama o MESMO ledger um
sinal por vez, com o estado persistido entre reinicios.

NAO muda nenhuma regra do detector/simulador nem a escolha da combinacao (isso continua vindo so de
`metrics.variant_table`/`choose_variant`, no custo FIXO de 0,15 SOL do handoff). So trades FECHADOS entram nesta
tabela e na contagem de posicoes abertas; cancelado/em_aberto ficam fora (limitacao documentada no pre-registro,
nao afeta a metrica oficial).
"""

from __future__ import annotations

import csv
import json
import logging
from pathlib import Path
from typing import Any, Optional

from ..bankroll import COLUMNS, HEADER_LABELS, BankrollLedger, BankrollParams
from ..detector import Event
from ..quotes import SOL_MINT
from .pipeline import (DEX_OF_LAUNCHPAD, IDADE_FAIXAS, QUEDA_MEDIDAS, Context, load_bars, primary_under_window,
                       tf_variant)

log = logging.getLogger("signals_mvp.pos_runner.parte_a.bankroll")

TFS = ("5m", "1h", "4h")
AJUSTES: tuple[tuple[str, str, str, str], ...] = tuple(
    (tf, medida, idade_label, variant)
    for tf in TFS for medida in QUEDA_MEDIDAS for idade_label in IDADE_FAIXAS for variant in ("A", "B")
)


def _fetch_token_name(ctx: Context, pool: str) -> Optional[str]:
    """Nome/simbolo do token via GeckoTerminal (so para exibicao na tabela; nunca filtro, nunca trava o relatorio)."""
    try:
        data = ctx.ohlcv.http.get_json(f"networks/{ctx.ohlcv.network}/pools/{pool}", {"include": "base_token"})
    except Exception:
        return None
    try:
        base_id = data["data"]["relationships"]["base_token"]["data"]["id"]
        for inc in data.get("included") or []:
            if inc.get("id") == base_id:
                sym = (inc.get("attributes") or {}).get("symbol")
                if sym:
                    return str(sym)
        name = ((data.get("data") or {}).get("attributes") or {}).get("name")
        return name.split(" / ")[0].strip() if name else None
    except Exception:
        return None


def _token_name(ctx: Context, pool: str) -> Optional[str]:
    if pool not in ctx._token_names:
        ctx._token_names[pool] = _fetch_token_name(ctx, pool)
    return ctx._token_names[pool]


def simulate_bankroll(ctx: Context, period: str, tf: str, medida: str, idade_label: str, variant: str,
                      params: BankrollParams = BankrollParams()) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    """Uma linha por entrada primaria FECHADA (ou "pulada") da combinacao `tf`/`medida`/`idade_label`/`variant`,
    em ordem cronologica real. "Primaria" aqui e reclassificada pela `idade_label` (ver `pipeline.primary_under_window`),
    nunca pelo `primary_flag` gravado na coleta (que so vale para a janela original 2-7)."""
    label = tf_variant(tf, medida)
    lo, hi = IDADE_FAIXAS[idade_label]
    prim_keys = primary_under_window(ctx, period, label, lo, hi)
    prim_rows = [r for r in ctx.store.setups(period, label, primary_only=False) if (r[0], r[1], r[2]) in prim_keys]
    closed = {(r[0], r[1], r[2]) for r in ctx.store.trades(period, "setup") if r[3] == variant and r[4] == "fechado"}
    grads = {g["pool"]: g for g in ctx.store.graduations(period)}
    sol_series = ctx.usd(SOL_MINT, period)
    sol_at = sol_series.at if sol_series else (lambda t: None)
    bars_cache: dict[str, list] = {}
    ledger = BankrollLedger(ctx.params, ctx.costs, params)

    rows: list[dict[str, Any]] = []
    for r in sorted((r for r in prim_rows if (r[0], r[1], r[2]) in closed), key=lambda r: r[4]):
        pool = r[0]
        g = grads.get(pool)
        if g is None:
            continue
        bars = bars_cache.setdefault(pool, load_bars(ctx, g, tf))
        if not bars:
            continue
        d = json.loads(r[5])
        ev = Event(kind=d["kind"], idx=d["idx"], ts=d["ts"], decision_ts=d["decision_ts"], data=d["data"])
        st1 = ctx.store.stage1(pool)
        l_bc = st1[3] if st1 else None
        dex = DEX_OF_LAUNCHPAD.get(g["launchpad"], "desconhecido")
        meta = {"token": _token_name(ctx, pool), "contrato": g["mint"], "launchpad": g["launchpad"], "tf": tf, "saida": variant}
        linha = ledger.step(bars, ev, variant, dex, l_bc, sol_at, meta)
        if linha is not None:
            rows.append(linha)
    return rows, ledger.state.resumo()


def _write_csv(path: Path, rows: list[dict[str, Any]], columns: tuple[str, ...] = COLUMNS,
              labels: dict[str, str] = HEADER_LABELS) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow([labels.get(c, c) for c in columns])
        for row in rows:
            w.writerow(["" if row.get(c) is None else row.get(c) for c in columns])


def xlsx_available() -> bool:
    try:
        import openpyxl  # noqa: F401
    except Exception:
        return False
    return True


def _write_sheet(wb: Any, name: str, rows: list[dict[str, Any]], columns: tuple[str, ...] = COLUMNS,
                labels: dict[str, str] = HEADER_LABELS) -> None:
    ws = wb.create_sheet(title=name[:31])
    ws.append([labels.get(c, c) for c in columns])
    for row in rows:
        ws.append([row.get(c) for c in columns])


# -- diagnostico (EMENDA 2026-10-16): NAO conta para o veredito, so para conferencia -----------------------------
DIAG_COLUMNS = ("tf", "pool", "n", "token", "contrato", "launchpad", "idade_dias", "n_entrada", "decision_ts")


def diagnostics(ctx: Context, period: str) -> dict[str, list[dict[str, Any]]]:
    """(a) entradas fora de QUALQUER janela de idade testada (idade < 1, > 10, ou desconhecida) e (b) reentradas
    (n_entrada > 1 no mesmo pool/tf) -- das 3 combinacoes "extremos" (a medida da queda nao importa aqui; idade e
    ordem de entrada sao as mesmas nas duas medidas). Nunca entram no veredito."""
    grads = {g["pool"]: g for g in ctx.store.graduations(period)}
    fora_idade: list[dict[str, Any]] = []
    reentradas: list[dict[str, Any]] = []
    lo_all = min(lo for lo, _hi in IDADE_FAIXAS.values())
    hi_all = max(hi for _lo, hi in IDADE_FAIXAS.values())
    for tf in TFS:
        for pool, tf_, n, _prim, ts, ev in ctx.store.setups(period, tf, primary_only=False):
            d = json.loads(ev)["data"]
            g = grads.get(pool)
            idade = d.get("idade_dias")
            base = {"tf": tf_, "pool": pool, "n": n, "token": _token_name(ctx, pool),
                   "contrato": g["mint"] if g else None, "launchpad": g["launchpad"] if g else None,
                   "idade_dias": idade, "n_entrada": d.get("n_entrada"), "decision_ts": ts}
            if idade is None or idade < lo_all or idade > hi_all:
                fora_idade.append(base)
            if (d.get("n_entrada") or 1) > 1:
                reentradas.append(base)
    return {"fora_de_qualquer_janela_de_idade": fora_idade, "reentradas_apos_1o_setup": reentradas}


def write_bankroll_reports(ctx: Context, docs: Path, periods: tuple[str, ...] = ("discovery", "holdout")) -> dict[str, Any]:
    """CSV por combinacao x periodo (sempre) + um XLSX com todas as abas (se `openpyxl` estiver instalado): 24
    abas de combinacao + 1 resumo + 2 de diagnostico (fora de qualquer janela de idade; reentradas)."""
    docs = Path(docs)
    out_dir = docs / "trades"
    resultado: dict[str, Any] = {"arquivos_csv": [], "resumo": {}}
    wb = None
    if xlsx_available():
        import openpyxl
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
    for period in periods:
        if not ctx.store.graduations(period):
            continue       # periodo ainda nao rodou: nenhuma tabela (nunca uma tabela vazia enganosa)
        for tf, medida, idade_label, variant in AJUSTES:
            rows, resumo = simulate_bankroll(ctx, period, tf, medida, idade_label, variant)
            name = f"{period}_{tf}_{medida}_{idade_label}_{variant}"
            path = out_dir / f"{name}.csv"
            _write_csv(path, rows)
            resultado["arquivos_csv"].append(str(path))
            resultado["resumo"][f"{period}/{tf}/{medida}/{idade_label}/{variant}"] = resumo
            if wb is not None:
                _write_sheet(wb, name, rows)
        diag = diagnostics(ctx, period)
        for diag_name, diag_rows in diag.items():
            path = out_dir / f"{period}_diag_{diag_name}.csv"
            _write_csv(path, diag_rows, DIAG_COLUMNS, {})
            resultado["arquivos_csv"].append(str(path))
            if wb is not None:
                _write_sheet(wb, f"{period}_diag_{diag_name}"[:31], diag_rows, DIAG_COLUMNS, {})
    if wb is not None:
        ws = wb.create_sheet(title="resumo", index=0)
        ws.append(["combinacao", "n_trades", "%_com_lucro", "lucro_liq_medio_%", "banca_final_usd", "maior_queda_usd", "maior_queda_%"])
        for k, s in resultado["resumo"].items():
            ws.append([k, s["n_trades"], s["pct_com_lucro"], s["lucro_liq_medio_pct"], s["banca_final_usd"], s["maior_queda_usd"], s["maior_queda_pct"]])
        xlsx_path = docs / "trades_parte_a.xlsx"
        wb.save(xlsx_path)
        resultado["arquivo_xlsx"] = str(xlsx_path)
    else:
        resultado["arquivo_xlsx"] = None
        log.warning("openpyxl indisponivel: so o CSV foi gerado (pip install openpyxl para o XLSX)")
    return resultado
